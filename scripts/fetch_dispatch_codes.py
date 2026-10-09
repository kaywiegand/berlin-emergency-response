"""Build the dispatch code seeds from the BF code table (AMPDS codes with emergency category).

Source:  Berliner Feuerwehr, BF-Open-Data, Datasets/Dispatchcodes/*.xlsx (CC BY 4.0), Stand in the file's Disclaimer sheet.
         The table covers ambulance service codes only (no fire, technical rescue, CBRN).
Output:  dbt/seeds/seed_dispatch_codes.csv        one row per complaint category (2 digits) with its name
         dbt/seeds/seed_dispatch_code_map.csv     one row per (category, level A-E/O): share of codes per emergency category RD1..RD5

A mission in Mission_Data carries only category and level (the first three of five code characters), so the emergency
category of a single mission is not determined. The map gives, per (category, level), the share of table codes that fall
into each category (unweighted by mission volume); it is a prior, not a lookup.
"""

from __future__ import annotations

import argparse
import csv
import pathlib
import re
from collections import Counter, defaultdict

import httpx
import openpyxl

BASE_DIR = pathlib.Path(__file__).resolve().parent.parent
DEFAULT_URL = (
    "https://raw.githubusercontent.com/Berliner-Feuerwehr/BF-Open-Data/main/Datasets/Dispatchcodes/"
    "Mediznische%20Codezuordnung%20Berliner%20Feuerwehr_20260521.xlsx"
)
RD_GROUPS = ["RD1", "RD2", "RD3", "RD4", "RD5"]


def load_rows(path: pathlib.Path) -> tuple[str, list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    stand = next((str(r[0]) for r in wb["Disclaimer"].iter_rows(values_only=True) if r[0] and str(r[0]).startswith("Stand")), "")
    rows = list(wb["AMPDS-Codes"].iter_rows(values_only=True))
    header = list(rows[0])
    out = []
    for r in rows[1:]:
        rec = dict(zip(header, r))
        if rec.get("Code"):
            out.append(rec)
    return stand, out


def build(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    names: dict[str, Counter] = defaultdict(Counter)
    shares: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for rec in rows:
        code = str(rec["Code"]).strip()
        if not re.fullmatch(r"\d{2}[A-EO]\d{2}[A-Z]?", code):
            continue
        cat, level = code[:2], code[2]
        if rec.get("Hauptbeschwerde_Text_Original"):
            names[cat][str(rec["Hauptbeschwerde_Text_Original"]).strip()] += 1
        m = re.match(r"(RD\d)", str(rec.get("Notfallkategorie") or ""))
        shares[(cat, level)][m.group(1) if m else "other"] += 1
    codes = [{"dispatchcode": cat, "dispatchcode_name": c.most_common(1)[0][0]} for cat, c in sorted(names.items())]
    mapping = []
    for (cat, level), c in sorted(shares.items()):
        total = sum(c.values())
        rec = {"dispatchcode": cat, "dispatch_level": level, "codes_in_table": total}
        for rd in RD_GROUPS + ["other"]:
            rec[f"share_{rd.lower()}"] = round(c[rd] / total, 4)
        rec["share_rd1_rd2"] = round((c["RD1"] + c["RD2"]) / total, 4)
        mapping.append(rec)
    return codes, mapping


def write_csv(path: pathlib.Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--url", default=DEFAULT_URL)
    p.add_argument("--xlsx", default=str(BASE_DIR / "data" / "raw" / "dispatch_codes" / "codes.xlsx"))
    p.add_argument("--seed-dir", default=str(BASE_DIR / "dbt" / "seeds"))
    p.add_argument("--refresh", action="store_true")
    args = p.parse_args()

    xlsx = pathlib.Path(args.xlsx)
    if args.refresh or not xlsx.exists():
        xlsx.parent.mkdir(parents=True, exist_ok=True)
        r = httpx.get(args.url, follow_redirects=True, timeout=120.0)
        r.raise_for_status()
        xlsx.write_bytes(r.content)
    stand, rows = load_rows(xlsx)
    codes, mapping = build(rows)
    seed_dir = pathlib.Path(args.seed_dir)
    write_csv(seed_dir / "seed_dispatch_codes.csv", codes)
    write_csv(seed_dir / "seed_dispatch_code_map.csv", mapping)
    print(f"{stand}: {len(rows)} code rows, {len(codes)} categories, {len(mapping)} (category, level) combinations")


if __name__ == "__main__":
    main()
